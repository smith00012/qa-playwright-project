import json
import re
from datetime import datetime
from html import escape
from pathlib import Path

from app.evidence.step_markers import draw_step_markers
from app.reporting.run_report import STYLES as REPORT_STYLES


GUIDE_DATA_FILENAME = "guide.json"
GUIDE_HTML_FILENAME = "guia.html"
GUIDE_DOCX_FILENAME = "guia.docx"
MAX_LABEL_LENGTH = 70
TYPED_KINDS = {"field", "password", "select"}
EDIT_HELP = (
    "Textos de la guía. Edite titulo, introduccion, el titulo y la nota de cada "
    "pantalla y el texto de cada paso (use **texto** para resaltar). Agregue a "
    "ocultar los números de paso que no deben aparecer. Se aplica la próxima vez "
    "que se genere la guía de este flujo."
)


def short_label(label: str) -> str:
    label = re.sub(r"\s*\*+\s*$", "", (label or "").strip())
    if len(label) <= MAX_LABEL_LENGTH:
        return label
    return label[:MAX_LABEL_LENGTH].rsplit(" ", 1)[0] + "…"


def default_step_text(step: dict) -> str:
    """Instrucción para una persona, sin los valores escritos durante la grabación."""

    label = short_label(step.get("label", ""))
    named = f"**{label}**" if label else ""
    kind = step.get("kind")
    if kind == "select":
        option = short_label(step.get("option", ""))
        if option:
            return f"En {named or 'la lista señalada'}, elige **{option}**."
        return f"Elige una opción en {named or 'la lista señalada'}."
    if kind == "password":
        return f"Escribe tu contraseña en {named or 'el campo señalado'}."
    if kind == "field":
        return f"Completa {named or 'el campo señalado'}."
    if kind == "check":
        return f"Marca {named or 'la casilla señalada'}."
    return f"Haz clic en {named or 'el elemento señalado'}."


def guide_steps(steps: list[dict]) -> list[dict]:
    """
    Pasos que se muestran: quita el clic que el grabador guarda antes de
    escribir en un campo, porque la instrucción de escribir ya lo incluye, y
    el aviso de cambio que el navegador da al salir de un campo ya escrito.
    """

    shown = []
    typed: set[str] = set()
    for step, following in zip(steps, steps[1:] + [None]):
        target = step.get("target")
        if (
            following
            and step.get("action_type") == "click"
            and target
            and target == following.get("target")
            and following.get("kind") in TYPED_KINDS
        ):
            continue
        if step.get("action_type") == "change" and step.get("kind") in {"field", "password"} and target in typed:
            continue
        if step.get("action_type") == "input" and target:
            typed.add(target)
        shown.append(step)
    return shown


def screen_title(title: str, url: str) -> str:
    if title:
        return title
    path = re.sub(r"^https?://[^/]+", "", url or "").strip("/")
    return path or url or "Pantalla"


def build_guide(data: dict, overrides: dict) -> dict:
    """
    Agrupa los pasos por pantalla (misma URL y título) y, dentro de cada
    pantalla, por captura. Numera los pasos de corrido; ese número es el
    que aparece marcado sobre la captura.
    """

    hidden = {int(number) for number in overrides.get("ocultar") or []}
    texts = overrides.get("pasos") or {}
    screen_overrides = overrides.get("pantallas") or {}
    screens: list[dict] = []
    number = 0

    for step in guide_steps(data.get("steps") or []):
        if step.get("step") in hidden:
            continue
        key = (step.get("url"), step.get("title"))
        if not screens or screens[-1]["key"] != key:
            screen_number = str(len(screens) + 1)
            custom = screen_overrides.get(screen_number) or {}
            screens.append({
                "key": key,
                "number": screen_number,
                "title": custom.get("titulo") or screen_title(step.get("title"), step.get("url")),
                "note": custom.get("nota") or "",
                "blocks": [],
            })
        blocks = screens[-1]["blocks"]
        if not blocks or blocks[-1]["shot"] != step.get("shot"):
            blocks.append({"shot": step.get("shot"), "steps": []})
        number += 1
        blocks[-1]["steps"].append({
            **step,
            "number": number,
            "text": texts.get(str(step.get("step"))) or default_step_text(step),
        })

    final = None
    if data.get("final_shot") is not None:
        custom = overrides.get("final") or {}
        shot = (data.get("shots") or [])[data["final_shot"]]
        final = {
            "shot": data["final_shot"],
            "title": custom.get("titulo") or screen_title(shot.get("title"), shot.get("url")),
            "note": custom.get("nota") or "Al terminar verás esta pantalla.",
        }

    return {"screens": screens, "final": final}


def overrides_template(flow_name: str, guide: dict) -> dict:
    return {
        "_ayuda": EDIT_HELP,
        "titulo": f"Guía de uso: {flow_name}",
        "introduccion": "",
        "pantallas": {
            screen["number"]: {"titulo": screen["title"], "nota": screen["note"]}
            for screen in guide["screens"]
        },
        "pasos": {
            str(step["step"]): step["text"]
            for screen in guide["screens"]
            for block in screen["blocks"]
            for step in block["steps"]
        },
        "ocultar": [],
        "final": {"titulo": "", "nota": ""},
    }


def segments(text: str) -> list[tuple[str, bool]]:
    """Divide un texto con **resaltados** en tramos (texto, negrita)."""

    return [(part, index % 2 == 1) for index, part in enumerate(text.split("**")) if part]


class UserGuideGenerator:
    """
    Genera guia.html y guia.docx de una corrida a partir de guide.json.

    Los textos editables de cada flujo viven en un archivo aparte (en
    data/guides/): si no existe se crea con los textos detectados, así la
    guía se puede regenerar con capturas nuevas sin perder lo redactado.
    """

    def generate(self, run_folder: Path, overrides_file: Path | None = None) -> Path:
        data = self._read(run_folder / GUIDE_DATA_FILENAME)
        result = self._read(run_folder / "result.json")
        flow_name = result.get("flow_name") or "Flujo"

        overrides = self._read(overrides_file) if overrides_file else {}
        guide = build_guide(data, overrides)
        if overrides_file and not overrides_file.exists():
            overrides_file.parent.mkdir(parents=True, exist_ok=True)
            with overrides_file.open("w", encoding="utf-8") as file:
                json.dump(overrides_template(flow_name, guide), file, ensure_ascii=False, indent=2)

        images = self._mark_shots(run_folder, data, guide)
        title = overrides.get("titulo") or f"Guía de uso: {flow_name}"
        intro = overrides.get("introduccion") or ""

        html_file = run_folder / GUIDE_HTML_FILENAME
        html_file.write_text(
            self.render_html(title, intro, guide, images, result),
            encoding="utf-8",
        )
        self.render_docx(run_folder / GUIDE_DOCX_FILENAME, title, intro, guide, images, run_folder)
        return html_file

    # ----------------------------------------------------------
    # Capturas
    # ----------------------------------------------------------

    @staticmethod
    def _mark_shots(run_folder: Path, data: dict, guide: dict) -> dict[int, str]:
        """Dibuja los números de los pasos sobre cada captura usada."""

        shots = data.get("shots") or []
        markers: dict[int, list] = {}
        for screen in guide["screens"]:
            for block in screen["blocks"]:
                markers.setdefault(block["shot"], []).extend(
                    (step["number"], step["box"]) for step in block["steps"] if step.get("box")
                )

        images: dict[int, str] = {}
        for index, shot_markers in markers.items():
            shot = shots[index]
            source = run_folder / shot["file"]
            if not source.exists():
                continue
            target = source.with_name(source.stem + "_pasos.png")
            draw_step_markers(source, target, shot_markers, shot.get("viewport_width"))
            images[index] = target.name

        final = guide.get("final")
        if final and (run_folder / shots[final["shot"]]["file"]).exists():
            images[final["shot"]] = shots[final["shot"]]["file"]
        return images

    # ----------------------------------------------------------
    # HTML
    # ----------------------------------------------------------

    def render_html(self, title: str, intro: str, guide: dict, images: dict, result: dict) -> str:
        index_items = "".join(
            f"<li><a href=\"#pantalla-{screen['number']}\">{escape(screen['title'])}</a></li>"
            for screen in guide["screens"]
        )
        sections = [self._html_screen(screen, images) for screen in guide["screens"]]
        final = guide.get("final")
        if final:
            sections.append(
                "<section id=\"final\"><h2>Resultado: " + escape(final["title"]) + "</h2>"
                f"<p>{self._html_text(final['note'])}</p>"
                f"{self._html_image(images.get(final['shot']), final['title'])}</section>"
            )
        intro_html = f"<p class=\"intro\">{self._html_text(intro)}</p>" if intro else ""
        index_html = (
            f"<nav><h2>Contenido</h2><ol>{index_items}</ol></nav>" if len(guide["screens"]) > 1 else ""
        )
        return (
            "<!doctype html>\n<html lang=\"es\">\n<head>\n<meta charset=\"utf-8\">\n"
            "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">\n"
            f"<title>{escape(title)}</title>\n<style>{REPORT_STYLES}{GUIDE_STYLES}</style>\n"
            "</head>\n<body>\n<main>\n"
            f"<header><p class=\"eyebrow\">Guía paso a paso</p><h1>{escape(title)}</h1>{intro_html}</header>\n"
            f"{index_html}\n" + "\n".join(sections) + "\n" + self._html_footer(result)
            + "\n</main>\n</body>\n</html>\n"
        )

    def _html_screen(self, screen: dict, images: dict) -> str:
        note = f"<p class=\"note\">{self._html_text(screen['note'])}</p>" if screen["note"] else ""
        blocks = []
        for block in screen["blocks"]:
            first = block["steps"][0]["number"]
            items = "".join(
                f"<li><span class=\"num\">{step['number']}</span><span>{self._html_text(step['text'])}</span></li>"
                for step in block["steps"]
            )
            blocks.append(
                "<div class=\"block\">"
                f"{self._html_image(images.get(block['shot']), screen['title'])}"
                f"<ol class=\"steps\" start=\"{first}\">{items}</ol></div>"
            )
        return (
            f"<section id=\"pantalla-{screen['number']}\">"
            f"<h2><span class=\"screen-num\">{screen['number']}</span>{escape(screen['title'])}</h2>"
            f"{note}{''.join(blocks)}</section>"
        )

    @staticmethod
    def _html_text(text: str) -> str:
        return "".join(
            f"<strong>{escape(part)}</strong>" if bold else escape(part)
            for part, bold in segments(text)
        )

    @staticmethod
    def _html_image(filename: str | None, alt: str) -> str:
        if not filename:
            return ""
        source = escape(filename)
        return f"<a href=\"{source}\" target=\"_blank\"><img class=\"shot\" src=\"{source}\" alt=\"{escape(alt)}\"></a>"

    @staticmethod
    def _html_footer(result: dict) -> str:
        notes = [
            f"Generada el {datetime.now().strftime('%d/%m/%Y %H:%M')} reproduciendo el flujo "
            f"«{result.get('flow_name', '')}» en {result.get('browser', '')} · {result.get('device', '')}.",
        ]
        if result.get("status") == "detenido_por_seguridad":
            notes.append(
                "La guía llega hasta el paso bloqueado "
                f"«{result.get('stopped_at') or ''}»: el freno de pago detuvo la reproducción en este ambiente."
            )
        elif result.get("status") == "fallido":
            notes.append(
                f"La reproducción falló en el paso {result.get('failed_step') or '?'}; la guía está incompleta."
            )
        items = "".join(f"<p>{escape(note)}</p>" for note in notes)
        return f"<footer>{items}</footer>"

    # ----------------------------------------------------------
    # Word
    # ----------------------------------------------------------

    def render_docx(
        self,
        target: Path,
        title: str,
        intro: str,
        guide: dict,
        images: dict,
        run_folder: Path,
    ) -> Path:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor

        document = Document()
        section = document.sections[0]
        section.left_margin = section.right_margin = Inches(0.8)
        document.styles["Normal"].font.name = "Aptos"
        document.styles["Normal"].font.size = Pt(10.5)
        image_width = section.page_width - section.left_margin - section.right_margin

        document.add_heading(title, level=0)
        if intro:
            self._docx_text(document.add_paragraph(), intro)

        for screen in guide["screens"]:
            document.add_heading(f"{screen['number']}. {screen['title']}", level=1)
            if screen["note"]:
                self._docx_text(document.add_paragraph(), screen["note"])
            for block in screen["blocks"]:
                image = images.get(block["shot"])
                if image:
                    document.add_picture(str(run_folder / image), width=image_width)
                for step in block["steps"]:
                    paragraph = document.add_paragraph()
                    number = paragraph.add_run(f"{step['number']}.  ")
                    number.bold = True
                    number.font.color.rgb = RGBColor(0xE6, 0x51, 0x00)
                    self._docx_text(paragraph, step["text"])

        final = guide.get("final")
        if final:
            document.add_heading(f"Resultado: {final['title']}", level=1)
            self._docx_text(document.add_paragraph(), final["note"])
            if images.get(final["shot"]):
                document.add_picture(str(run_folder / images[final["shot"]]), width=image_width)

        document.save(str(target))
        return target

    @staticmethod
    def _docx_text(paragraph, text: str) -> None:
        for part, bold in segments(text):
            paragraph.add_run(part).bold = bold

    @staticmethod
    def _read(filepath: Path) -> dict:
        if not filepath.exists():
            return {}
        with filepath.open("r", encoding="utf-8") as file:
            return json.load(file)


GUIDE_STYLES = """
main { max-width: 960px; }
header .intro { margin: 8px 0 0; color: #dbe6e2; max-width: 70ch; }
nav { background: var(--panel); border: 1px solid var(--border); border-radius: 10px;
  padding: 16px 20px; margin-bottom: 16px; }
nav ol { margin: 0; padding-left: 20px; columns: 2 260px; }
h2 { display: flex; align-items: center; gap: 10px; font-size: 19px; }
.screen-num { display: inline-grid; place-items: center; width: 30px; height: 30px;
  border-radius: 50%; background: var(--brand); color: #fff; font-size: 14px; flex: none; }
.block + .block { margin-top: 20px; padding-top: 20px; border-top: 1px dashed var(--border); }
ol.steps { list-style: none; margin: 14px 0 0; padding: 0; }
ol.steps li { display: flex; gap: 12px; align-items: baseline; padding: 6px 0; font-size: 16px; }
ol.steps .num { display: inline-grid; place-items: center; min-width: 26px; height: 26px;
  border-radius: 50%; background: #e65100; color: #fff; font-size: 13px; font-weight: 700; flex: none; }
img.shot { margin-top: 0; }
footer { color: var(--muted); font-size: 13px; padding: 8px 4px; }
footer p { margin: 4px 0; }
"""
